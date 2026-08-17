import asyncio
import hashlib
import json
import os

from io import BytesIO
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from runtime import DocumentRuntimeError, convert_bytes, safe_filename
from main import convert_document


class StreamRequest:
    def __init__(self, chunks: list[bytes], mime_type: str = "text/plain") -> None:
        self.headers = {"content-type": mime_type}
        self.chunks = chunks
        self.yielded = 0

    async def stream(self):
        for chunk in self.chunks:
            self.yielded += 1
            yield chunk


def pdf_fixture(text: str) -> bytes:
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    document = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(document))
        document.extend(f"{index} 0 obj\n".encode("ascii"))
        document.extend(value)
        document.extend(b"\nendobj\n")
    xref = len(document)
    document.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    document.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    document.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    return bytes(document)


def docx_fixture(text: str) -> bytes:
    body = BytesIO()
    with ZipFile(body, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">
  <Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>
  <Default Extension=\"xml\" ContentType=\"application/xml\"/>
  <Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">
  <Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/>
</Relationships>""",
        )
        archive.writestr(
            "word/document.xml",
            f"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">
  <w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p><w:sectPr/></w:body>
</w:document>""",
        )
    return body.getvalue()


def pptx_fixture(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = text
    body = BytesIO()
    presentation.save(body)
    return body.getvalue()


def xlsx_fixture(text: str) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Brief"
    workbook.active["A1"] = text
    body = BytesIO()
    workbook.save(body)
    return body.getvalue()


class DocumentRuntimeTests(unittest.TestCase):
    def test_plain_text_uses_bounded_stream_conversion(self) -> None:
        result = convert_bytes(body=b"# Company\n\nTrusted local document.", mime_type="text/plain", filename="company.txt")
        self.assertIn("Company", result.markdown)
        self.assertEqual(result.converter, "markitdown")

    def test_rejects_url_like_or_mismatched_filename(self) -> None:
        with self.assertRaises(DocumentRuntimeError):
            safe_filename("https://example.invalid/file.txt", "text/plain")
        with self.assertRaises(DocumentRuntimeError):
            safe_filename("company.pdf", "text/plain")

    def test_rejects_unknown_mime_without_network_or_path_access(self) -> None:
        with self.assertRaises(DocumentRuntimeError) as caught:
            convert_bytes(body=b"data", mime_type="application/octet-stream", filename="input.bin")
        self.assertEqual(caught.exception.code, "DOCUMENT_UNSUPPORTED")

    def test_supported_office_and_pdf_fixtures_convert_from_memory(self) -> None:
        fixtures = [
            ("application/pdf", "brief.pdf", pdf_fixture("C10 PDF fixture"), "C10 PDF fixture"),
            (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "brief.docx",
                docx_fixture("C10 DOCX fixture"),
                "C10 DOCX fixture",
            ),
            (
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "brief.pptx",
                pptx_fixture("C10 PPTX fixture"),
                "C10 PPTX fixture",
            ),
            (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "brief.xlsx",
                xlsx_fixture("C10 XLSX fixture"),
                "C10 XLSX fixture",
            ),
        ]

        for mime_type, filename, body, expected in fixtures:
            with self.subTest(filename=filename):
                result = convert_bytes(body=body, mime_type=mime_type, filename=filename)
                self.assertIn(expected, result.markdown)
                self.assertEqual(result.converter, "markitdown")

    def test_conversion_invokes_only_the_bounded_stream_entrypoint(self) -> None:
        converted = SimpleNamespace(text_content="# Stream only")
        converter = SimpleNamespace(convert_stream=lambda stream, stream_info: converted)
        with patch("runtime.MarkItDown", return_value=converter) as markitdown:
            result = convert_bytes(body=b"# Stream only", mime_type="text/markdown", filename="brief.md")

        self.assertEqual(result.markdown, "# Stream only")
        markitdown.assert_called_once_with(enable_plugins=False)

    def test_http_handler_rejects_oversized_stream_before_aggregating_remaining_chunks(self) -> None:
        request = StreamRequest([b"abc", b"de", b"never-read"])
        with patch.dict(os.environ, {"DOCUMENT_RUNTIME_TOKEN": "runtime-test-token"}, clear=False), patch("main.MAX_INPUT_BYTES", 4):
            response = asyncio.run(
                convert_document(
                    request=request,
                    authorization="Bearer runtime-test-token",
                    x_document_conversion_id="dcv_01J00000000000000000000000",
                    x_source_filename="brief.txt",
                    x_source_sha256=hashlib.sha256(b"abcde").hexdigest(),
                )
            )

        self.assertEqual(response.status_code, 413)
        self.assertEqual(json.loads(response.body), {"error": {"code": "DOCUMENT_UNSUPPORTED", "retryable": False}})
        self.assertEqual(request.yielded, 2)


if __name__ == "__main__":
    unittest.main()
